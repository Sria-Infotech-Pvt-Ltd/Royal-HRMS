"""
Helpers for generating downloadable CSV and XLSX file content.

Used exclusively by bulk-import sample template endpoints.
openpyxl is already declared in requirements.txt (openpyxl==3.1.5).
"""
from __future__ import annotations

import csv
import io


_XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_CSV_MIME  = 'text/csv; charset=utf-8'


def build_sample_csv(headers: list, rows: list) -> bytes:
    """Return UTF-8 BOM CSV bytes so Excel opens without encoding prompts."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode('utf-8-sig')


def build_sample_xlsx(headers: list, rows: list, sheet_name: str = 'Template') -> bytes:
    """Return XLSX bytes with a styled header row and sample data rows."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='4472C4')
    center      = Alignment(horizontal='center', vertical='center', wrap_text=False)

    for col, text in enumerate(headers, 1):
        cell            = ws.cell(row=1, column=col, value=text)
        cell.font       = header_font
        cell.fill       = header_fill
        cell.alignment  = center

    for row_idx, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col, value=value)

    # Auto-fit column widths (capped at 40)
    for col, text in enumerate(headers, 1):
        letter  = get_column_letter(col)
        max_len = len(text)
        for row in rows:
            if col - 1 < len(row) and row[col - 1] is not None:
                max_len = max(max_len, len(str(row[col - 1])))
        ws.column_dimensions[letter].width = min(max_len + 4, 40)

    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
