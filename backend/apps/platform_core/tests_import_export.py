"""Phase 2 Task I — import/export framework tests. Uses the
"Vehicle Allocation" acceptance-criteria scenario: create records via
import, export them, roll back a create-mode import."""
from __future__ import annotations

import io

import openpyxl
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from apps.platform_core import services_custom_objects as custom_objects
from apps.platform_core import services_import_export as import_export
from apps.platform_core.models import EntityDefinition, FieldDefinition, ImportJob


def _xlsx_bytes(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportExportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.entity = EntityDefinition.objects.create(
            code='vehicle_allocation', label='Vehicle Allocation', kind=EntityDefinition.KIND_CUSTOM,
            attributes_column='data', title_template='{vehicle_number}',
        )
        self.vehicle_field = FieldDefinition.objects.create(
            entity=self.entity, code='vehicle_number', label='Vehicle Number', data_type='text',
            status=FieldDefinition.STATUS_PUBLISHED, required=True, is_unique=True, is_filterable=True, order=0,
        )
        self.driver_field = FieldDefinition.objects.create(
            entity=self.entity, code='driver_name', label='Driver Name', data_type='text',
            status=FieldDefinition.STATUS_PUBLISHED, order=1,
        )

    def _make_job(self, rows, *, mode=ImportJob.MODE_CREATE, key_field=''):
        content = _xlsx_bytes(rows)
        upload = SimpleUploadedFile('import.xlsx', content, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        return ImportJob.objects.create(entity=self.entity, mode=mode, key_field=key_field, file=upload)

    def test_generate_import_template_has_one_column_per_importable_field(self):
        wb = import_export.generate_import_template(self.entity)
        ws = wb['Data']
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, ['vehicle_number', 'driver_name'])
        self.assertIn('Instructions', wb.sheetnames)

    def test_validate_job_separates_valid_and_invalid_rows(self):
        job = self._make_job([
            ['vehicle_number', 'driver_name'],
            ['TS09AB1234', 'Ravi'],
            ['', 'Kiran'],  # missing required vehicle_number
        ])
        valid_rows, errors_by_row = import_export.validate_job(job)
        self.assertEqual(len(valid_rows), 1)
        self.assertIn(3, errors_by_row)  # header=row1, first data row=row2, second=row3
        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.STATUS_VALIDATED)
        self.assertEqual(job.valid_rows, 1)
        self.assertEqual(job.invalid_rows, 1)

    def test_commit_job_creates_records_and_tracks_ids_for_rollback(self):
        job = self._make_job([
            ['vehicle_number', 'driver_name'],
            ['TS09AB1234', 'Ravi'],
            ['TS09CD5678', 'Kiran'],
        ])
        counts = import_export.commit_job(job)
        self.assertEqual(counts, {'created': 2, 'updated': 0, 'skipped': 0})
        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.STATUS_COMPLETED)
        self.assertEqual(len(job.created_record_ids), 2)
        self.assertEqual(custom_objects.list_records(self.entity).count(), 2)

    def test_commit_job_skips_invalid_rows_by_default(self):
        job = self._make_job([
            ['vehicle_number', 'driver_name'],
            ['TS09AB1234', 'Ravi'],
            ['', 'Kiran'],
        ])
        counts = import_export.commit_job(job)
        self.assertEqual(counts['created'], 1)
        self.assertEqual(counts['skipped'], 1)

    def test_rollback_soft_deletes_only_created_records(self):
        job = self._make_job([
            ['vehicle_number', 'driver_name'],
            ['TS09AB1234', 'Ravi'],
        ])
        import_export.commit_job(job)
        self.assertEqual(custom_objects.list_records(self.entity).count(), 1)

        removed = import_export.rollback_job(job)
        self.assertEqual(removed, 1)
        self.assertEqual(custom_objects.list_records(self.entity).count(), 0)
        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.STATUS_ROLLED_BACK)

    def test_rollback_rejected_for_update_mode_job(self):
        job = self._make_job(
            [['vehicle_number', 'driver_name'], ['TS09AB1234', 'Ravi']],
            mode=ImportJob.MODE_UPDATE, key_field='vehicle_number',
        )
        job.status = ImportJob.STATUS_COMPLETED
        job.save(update_fields=['status'])
        with self.assertRaises(import_export.ImportExportError):
            import_export.rollback_job(job)

    def test_upsert_mode_updates_existing_and_creates_new(self):
        custom_objects.create_record(self.entity, {'vehicle_number': 'TS09AB1234', 'driver_name': 'OldDriver'})
        job = self._make_job(
            [['vehicle_number', 'driver_name'], ['TS09AB1234', 'NewDriver'], ['TS09CD5678', 'Kiran']],
            mode=ImportJob.MODE_UPSERT, key_field='vehicle_number',
        )
        counts = import_export.commit_job(job)
        self.assertEqual(counts, {'created': 1, 'updated': 1, 'skipped': 0})
        updated_record = custom_objects.list_records(self.entity).get(data__vehicle_number='TS09AB1234')
        self.assertEqual(updated_record.data['driver_name'], 'NewDriver')

    def test_export_escapes_formula_injection(self):
        content = import_export.export_rows_to_xlsx(
            self.entity, [{'vehicle_number': '=cmd|/c calc', 'driver_name': 'Ravi'}],
            columns=['vehicle_number', 'driver_name'],
        )
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        cell_value = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0][0]
        self.assertTrue(cell_value.startswith("'="))

    def test_import_row_limit_enforced(self):
        rows = [['vehicle_number', 'driver_name']] + [[f'TS{i:04d}', 'X'] for i in range(import_export.MAX_IMPORT_ROWS + 1)]
        job = self._make_job(rows)
        with self.assertRaises(import_export.ImportExportError):
            import_export.validate_job(job)
