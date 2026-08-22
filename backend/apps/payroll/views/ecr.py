import io
import logging
from decimal import Decimal, ROUND_HALF_UP

from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated

from core.responses import error
from core.permissions import has_perm as _has_perm
from apps.payroll.models import PayrollCycle, EmployeePayslip, PayrollSettings, BranchPayrollConfig
from apps.branch.models import Branch

logger = logging.getLogger(__name__)

_ZERO = Decimal('0')
_PF_FALLBACK_CEILING = Decimal('15000.00')


def _round_int(value: Decimal) -> int:
    """Round a Decimal to the nearest integer using standard half-up rounding."""
    return int(value.quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def _get_authorized_cycle(request, cycle_pk, resource_label: str = 'ECR file'):
    """
    Resolve + authorize a payroll cycle for a statutory export — the
    permission, existence, branch-scoping, and status checks shared by
    every such export (ECR Excel/PDF/text, and the ESIC Challan in
    esic.py, which imports this rather than duplicating it), extracted so
    they can't drift between formats or file types.

    resource_label customizes the permission/status error text only (e.g.
    'ESIC Challan') — every other check and message stays identical
    regardless of caller.

    Returns (cycle, None) on success, or (None, error_response) if access
    should be denied — same checks, same order, same messages the Excel
    download already used before this was extracted.
    """
    if not (_has_perm(request.user, 'payroll.view') or _has_perm(request.user, 'payroll.edit')):
        return None, error(f'You do not have permission to download the {resource_label}.', http_status=403)

    try:
        cycle = PayrollCycle.objects.select_related('branch').get(pk=cycle_pk)
    except PayrollCycle.DoesNotExist:
        return None, error('Payroll cycle not found.', http_status=404)

    # Branch-scoped access: a non-superuser can only download for their own branch.
    # Users with no branch assigned (global admins) may access any cycle.
    if not getattr(request.user, 'is_superuser', False):
        user_branch_id = getattr(request.user, 'branch_id', None)
        if cycle.branch_id and user_branch_id and cycle.branch_id != user_branch_id:
            return None, error('You do not have access to this payroll cycle.', http_status=403)

    if cycle.status not in ('payslips_generated', 'query_window_open', 'paid', 'closed'):
        return None, error(
            f'{resource_label} is only available after payslips have been generated.',
            http_status=400,
        )

    return cycle, None


def _ecr_filename(cycle: PayrollCycle, extension: str) -> str:
    period = cycle.cycle_start.strftime('%b_%Y') if cycle.cycle_start else 'cycle'
    try:
        branch_name = cycle.branch.branch_name if cycle.branch_id else 'All'
    except Exception:
        branch_name = 'All'
    branch = branch_name.replace(' ', '_')
    return f'ECR_{period}_{branch}.{extension}'


def _resolve_pf_wage_base(ps: EmployeePayslip, branch_config_by_name: dict) -> Decimal:
    """
    The wage base PF was actually calculated on for this payslip. Prefers
    the value stored at processing time (pf_wage_base — exact, by
    construction). Payslips computed before that field existed fall back to
    recomputing from the employee's OWN branch config (not the unrelated
    company-wide PayrollSettings.edli_wage_ceiling this function used to
    read) — same basic_only / basic_plus_allowances rule, same ceiling,
    that apps/payroll/views/cycles.py._compute_employee_payslip applied.
    This is a best-effort reconstruction for legacy rows only; it's exactly
    why validate_ecr() (services_ecr_filing.py) recomputes independently
    and flags a mismatch rather than trusting this silently.
    """
    if ps.pf_wage_base is not None:
        return ps.pf_wage_base

    basic = ps.basic or _ZERO
    branch_config = branch_config_by_name.get(getattr(ps.employee, 'branch', None))
    if branch_config is None:
        return min(basic, _PF_FALLBACK_CEILING)

    if branch_config.pf_wage_basis == BranchPayrollConfig.PF_WAGE_BASIC_PLUS_ALLOWANCES:
        uncapped = basic + (ps.special_allowance or _ZERO) + _other_earnings_total(ps)
    else:
        uncapped = basic
    return min(uncapped, branch_config.pf_wage_ceiling)


def _other_earnings_total(ps: EmployeePayslip) -> Decimal:
    return sum((Decimal(str(v)) for v in (ps.other_earnings or {}).values()), _ZERO)


def _compute_ecr_rows(cycle: PayrollCycle, settings: PayrollSettings) -> list:
    """Return one dict per employee payslip, with all ECR columns computed."""
    eps_rate = settings.eps_rate / 100
    edli_wage_ceil = settings.edli_wage_ceiling
    rows = []

    payslips = list(
        EmployeePayslip.objects
        .filter(cycle=cycle)
        .select_related('employee', 'employee__profile')
        .order_by('employee__full_name')
    )

    # Only needed as a fallback for payslips predating pf_wage_base — see
    # _resolve_pf_wage_base's docstring.
    branch_names = {getattr(ps.employee, 'branch', None) for ps in payslips}
    branch_config_by_name = {
        b.branch_name: b.payroll_config
        for b in Branch.objects.filter(branch_name__in=branch_names).select_related('payroll_config')
        if hasattr(b, 'payroll_config')
    }

    for idx, ps in enumerate(payslips, start=1):
        emp     = ps.employee
        profile = getattr(emp, 'profile', None)

        basic        = ps.basic or _ZERO
        gross        = ps.gross_earnings or _ZERO
        pf_wage      = _resolve_pf_wage_base(ps, branch_config_by_name)
        epf_contrib  = ps.pf_employee or _ZERO                   # employee 12%
        eps_contrib  = (pf_wage * eps_rate).quantize(Decimal('1'), rounding=ROUND_HALF_UP)
        epf_diff     = max(_ZERO, epf_contrib - eps_contrib)
        ncp_days     = _round_int(ps.lop_days or _ZERO)

        rows.append({
            'sr_no':              idx,
            'employee_name':      emp.full_name or '',
            'uan':                (profile.uan_number        if profile else '') or '',
            'name_as_per_aadhar': (profile.name_as_per_aadhar if profile else '') or '',
            'gross_wages':        _round_int(gross),
            'basic_wages':        _round_int(basic),
            'pension_wages':      _round_int(pf_wage),
            'edli_wages':         _round_int(min(basic, edli_wage_ceil)),
            'epf_contribution':   _round_int(epf_contrib),
            'eps_contribution':   int(eps_contrib),
            'epf_eps_difference': _round_int(epf_diff),
            'ncp_days':           ncp_days,
            # Carried through for validate_ecr()'s reconciliation checks —
            # not part of the EPFO column layout, harmless extra dict key
            # for the xlsx/pdf/text builders below (they only read named keys).
            'pf_employer':        ps.pf_employer or _ZERO,
        })

    return rows


def _build_ecr_xlsx(cycle: PayrollCycle, rows: list) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ECR'

    # ── Header ─────────────────────────────────────────────────────────────────
    header_fill   = PatternFill('solid', fgColor='1A3A6E')
    header_font   = Font(bold=True, color='FFFFFF', size=10)
    center_align  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin          = Side(border_style='thin', color='CCCCCC')
    cell_border   = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        'Sr No', 'Employee Name', 'UAN',
        'Employee Name as per Aadhar',
        'Gross Wages', 'Basic Wages', 'Pension Wages', 'EDLI Wages',
        'EPF Contribution', 'EPS Contribution', 'EPF and EPS Difference', 'NCP',
    ]
    col_widths = [6, 28, 14, 30, 12, 12, 12, 12, 14, 14, 20, 6]

    ws.row_dimensions[1].height = 36
    for col, (title, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center_align
        cell.border    = cell_border
        ws.column_dimensions[cell.column_letter].width = width

    # ── Data rows ──────────────────────────────────────────────────────────────
    num_font    = Font(size=10)
    num_align   = Alignment(horizontal='right', vertical='center')
    text_align  = Alignment(horizontal='left',  vertical='center')

    totals = {k: 0 for k in [
        'gross_wages', 'basic_wages', 'pension_wages', 'edli_wages',
        'epf_contribution', 'eps_contribution', 'epf_eps_difference', 'ncp_days',
    ]}

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row['sr_no'], row['employee_name'], row['uan'],
            row['name_as_per_aadhar'],
            row['gross_wages'], row['basic_wages'], row['pension_wages'], row['edli_wages'],
            row['epf_contribution'], row['eps_contribution'], row['epf_eps_difference'], row['ncp_days'],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border    = cell_border
            cell.font      = num_font
            cell.alignment = num_align if isinstance(value, (int, float, Decimal)) else text_align

        for key in totals:
            totals[key] += row[key]

    # ── Totals row ─────────────────────────────────────────────────────────────
    total_row = len(rows) + 2
    total_fill = PatternFill('solid', fgColor='FFF9C4')
    total_font = Font(bold=True, size=10)

    total_values = [
        '', 'TOTAL', '', '',
        totals['gross_wages'], totals['basic_wages'], totals['pension_wages'], totals['edli_wages'],
        totals['epf_contribution'], totals['eps_contribution'], totals['epf_eps_difference'], totals['ncp_days'],
    ]
    for col, value in enumerate(total_values, start=1):
        cell = ws.cell(row=total_row, column=col, value=value)
        cell.font      = total_font
        cell.fill      = total_fill
        cell.border    = cell_border
        cell.alignment = num_align if isinstance(value, (int, float)) else text_align

    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_ecr_pdf(cycle: PayrollCycle, rows: list) -> bytes:
    """Same ECR columns/totals as _build_ecr_xlsx above, laid out as a PDF table."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    period_label = cycle.cycle_start.strftime('%B %Y') if cycle.cycle_start else ''
    try:
        branch_label = cycle.branch.branch_name if cycle.branch_id else 'All Branches'
    except Exception:
        branch_label = 'All Branches'

    headers = [
        'Sr No', 'Employee Name', 'UAN', 'Employee Name\nas per Aadhar',
        'Gross\nWages', 'Basic\nWages', 'Pension\nWages', 'EDLI\nWages',
        'EPF\nContribution', 'EPS\nContribution', 'EPF-EPS\nDifference', 'NCP',
    ]
    table_data = [headers]

    totals = {k: 0 for k in [
        'gross_wages', 'basic_wages', 'pension_wages', 'edli_wages',
        'epf_contribution', 'eps_contribution', 'epf_eps_difference', 'ncp_days',
    ]}
    for row in rows:
        table_data.append([
            row['sr_no'], row['employee_name'], row['uan'], row['name_as_per_aadhar'],
            row['gross_wages'], row['basic_wages'], row['pension_wages'], row['edli_wages'],
            row['epf_contribution'], row['eps_contribution'], row['epf_eps_difference'], row['ncp_days'],
        ])
        for key in totals:
            totals[key] += row[key]

    table_data.append([
        '', 'TOTAL', '', '',
        totals['gross_wages'], totals['basic_wages'], totals['pension_wages'], totals['edli_wages'],
        totals['epf_contribution'], totals['eps_contribution'], totals['epf_eps_difference'], totals['ncp_days'],
    ])

    table = Table(table_data, repeatRows=1)
    last_row = len(table_data) - 1
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A3A6E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTNAME', (0, last_row), (-1, last_row), 'Helvetica-Bold'),
        ('BACKGROUND', (0, last_row), (-1, last_row), colors.HexColor('#FFF9C4')),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('ALIGN', (4, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, last_row - 1), [colors.white, colors.HexColor('#F7F9FC')]),
    ]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm, bottomMargin=12 * mm,
    )
    doc.build([
        Paragraph(f'ECR (Electronic Challan-cum-Return) — {period_label}', styles['Heading2']),
        Paragraph(f'Branch: {branch_label}', styles['Normal']),
        Spacer(1, 10),
        table,
    ])
    return buf.getvalue()


def _build_ecr_text(rows: list) -> str:
    """
    EPFO Unified Portal ECR text file — one '#~#'-delimited line per employee,
    in the field order the portal's ECR upload expects: UAN, Member Name,
    Gross Wages, EPF Wages, EPS Wages, EDLI Wages, EPF Contribution Remitted,
    EPS Contribution Remitted, EPF-EPS Difference Remitted, NCP Days, Refund
    of Advances. This is the publicly documented field layout; EPFO has
    revised the exact spec before (this codebase's own ECR history — see
    migration comments — already shows drift caused by manual DB patches
    elsewhere), so validate the very first real upload against the current
    Unified Portal spec before relying on it for a live filing.
    """
    lines = []
    for row in rows:
        # EPFO's Unified Portal expects the name exactly as registered against
        # the UAN (Aadhaar-linked KYC) — unlike the xlsx/pdf exports (which
        # show employee_name and name_as_per_aadhar as two separate columns
        # for HR's own cross-checking), the actual upload file must use the
        # Aadhaar name; a mismatch against full_name (nicknames, casing,
        # middle-name differences) is a real cause of portal upload rejection.
        # Falls back to employee_name only when name_as_per_aadhar was never
        # captured, so the row isn't left blank.
        member_name = row['name_as_per_aadhar'] or row['employee_name']
        lines.append('#~#'.join(str(v) for v in [
            row['uan'],
            member_name,
            row['gross_wages'],
            row['basic_wages'],
            row['pension_wages'],
            row['edli_wages'],
            row['epf_contribution'],
            row['eps_contribution'],
            row['epf_eps_difference'],
            row['ncp_days'],
            0,  # Refund of Advances — not tracked in this system; always 0
        ]) + '#~#')
    return '\r\n'.join(lines)


class CycleECRDownloadView(APIView):
    """GET /payroll/cycles/<pk>/ecr/  — download ECR Excel for a payroll cycle."""

    permission_classes = [IsAuthenticated]

    def get(self, request, cycle_pk):
        cycle, err = _get_authorized_cycle(request, cycle_pk)
        if err:
            return err

        settings_obj = PayrollSettings.objects.first()
        if settings_obj is None:
            settings_obj = PayrollSettings()

        rows = _compute_ecr_rows(cycle, settings_obj)
        if not rows:
            return error('No payslips found for this cycle.', http_status=404)

        try:
            xlsx_bytes = _build_ecr_xlsx(cycle, rows)
        except Exception as exc:
            logger.error('ECR xlsx generation failed for cycle %s: %s', cycle_pk, exc)
            return error('Failed to generate ECR file. Please try again.', http_status=500)

        filename = _ecr_filename(cycle, 'xlsx')
        response = HttpResponse(
            xlsx_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        logger.info('ECR downloaded for cycle %s by %s', cycle_pk, request.user.email)
        return response


class CycleECRPdfDownloadView(APIView):
    """GET /payroll/cycles/<pk>/ecr/pdf/  — download ECR as PDF for a payroll cycle."""

    permission_classes = [IsAuthenticated]

    def get(self, request, cycle_pk):
        cycle, err = _get_authorized_cycle(request, cycle_pk)
        if err:
            return err

        settings_obj = PayrollSettings.objects.first()
        if settings_obj is None:
            settings_obj = PayrollSettings()

        rows = _compute_ecr_rows(cycle, settings_obj)
        if not rows:
            return error('No payslips found for this cycle.', http_status=404)

        try:
            pdf_bytes = _build_ecr_pdf(cycle, rows)
        except Exception as exc:
            logger.error('ECR pdf generation failed for cycle %s: %s', cycle_pk, exc)
            return error('Failed to generate ECR PDF file. Please try again.', http_status=500)

        filename = _ecr_filename(cycle, 'pdf')
        response = HttpResponse(pdf_bytes, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        logger.info('ECR PDF downloaded for cycle %s by %s', cycle_pk, request.user.email)
        return response


class CycleECRTextDownloadView(APIView):
    """GET /payroll/cycles/<pk>/ecr-text/ — EPFO Unified Portal ECR upload text file."""

    permission_classes = [IsAuthenticated]

    def get(self, request, cycle_pk):
        cycle, err = _get_authorized_cycle(request, cycle_pk)
        if err:
            return err

        settings_obj = PayrollSettings.objects.first() or PayrollSettings()
        rows = _compute_ecr_rows(cycle, settings_obj)
        if not rows:
            return error('No payslips found for this cycle.', http_status=404)

        missing_uan = [r['employee_name'] for r in rows if not r['uan']]
        if missing_uan:
            return error(
                'The following employees have no UAN on file — add it before filing: '
                + ', '.join(missing_uan)
            )

        text = _build_ecr_text(rows)
        filename = _ecr_filename(cycle, 'txt')

        response = HttpResponse(text, content_type='text/plain')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        logger.info('ECR text file downloaded for cycle %s by %s', cycle_pk, request.user.email)
        return response
