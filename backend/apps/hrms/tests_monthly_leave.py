"""
Monthly Leave Management (Settings → Leave Management).

Covers set_monthly_credit() (difference-only application, edits, zero,
idempotent re-apply, validation, history kept), the dashboard cache drop,
import preview/confirm (templates, errors, duplicates, re-import, partial
failure, concurrency lock), manual entry, permissions and branch scope,
and that annual/opening/used/carry-forward values are untouched.
"""
from __future__ import annotations

import csv
import io
import threading
from decimal import Decimal

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, connections
from django.test import TestCase, TransactionTestCase
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.models import AuditLog
from apps.hrms.models import CREDIT_TXN_MONTHLY, LeaveBalance, LeaveCreditTransaction
from apps.hrms.services_leave_credit import (
    LeaveCreditError, monthly_credit_value, reverse_transaction, set_monthly_credit, yearly_summary,
)
from apps.hrms.tests_leave_credits import PASSWORD, YEAR, _Fixtures, _key
from config.test_runner import TEST_COMPANY_CODE

MONTH = 3


def _csv(rows, header=('Employee Code', 'Leave Type', 'Monthly Leave Credit'), name='monthly.csv'):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    return SimpleUploadedFile(name, buf.getvalue().encode('utf-8'), content_type='text/csv')


def _xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Employee Code', 'Leave Type', 'Monthly Leave Credit'])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile('monthly.xlsx', buf.getvalue(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class MonthlyCreditServiceTests(_Fixtures, TestCase):
    def setUp(self):
        self.make_fixtures()

    def set(self, value, **kw):
        kw.setdefault('employee', self.employee)
        kw.setdefault('leave_type', 'casual')
        kw.setdefault('year', YEAR)
        kw.setdefault('month', MONTH)
        kw.setdefault('actor', self.hr_hyd)
        return set_monthly_credit(value=value, **kw)

    def total(self):
        self.balance.refresh_from_db()
        return self.balance.total_days

    def test_add_edit_reapply_and_zero(self):
        self.assertEqual(self.set('1.5')['status'], 'created')
        self.assertEqual(self.total(), Decimal('13.5'))

        result = self.set('2')
        self.assertEqual((result['status'], result['delta']), ('updated', Decimal('0.5')))
        self.assertEqual(self.total(), Decimal('14'))

        self.assertEqual(self.set('2')['status'], 'unchanged')   # same value: no change, no row
        self.assertEqual(self.total(), Decimal('14'))

        self.assertEqual(self.set('0')['delta'], Decimal('-2'))  # 0 reverses the month's credit
        self.assertEqual(self.total(), Decimal('12'))
        self.assertEqual(monthly_credit_value(self.employee, 'casual', YEAR, MONTH), Decimal('0'))
        # history is append-only: three monthly rows remain (+1.5, +0.5, -2)
        rows = LeaveCreditTransaction.objects.filter(txn_type=CREDIT_TXN_MONTHLY).order_by('created_at')
        self.assertEqual([r.amount for r in rows], [Decimal('1.5'), Decimal('0.5'), Decimal('-2')])
        self.assertTrue(all(r.month == MONTH for r in rows))
        self.assertEqual(AuditLog.objects.filter(action='leave_monthly_credit_set').count(), 3)

    def test_months_and_leave_types_are_independent(self):
        self.set('1')
        self.set('2', month=MONTH + 1)
        self.assertEqual(monthly_credit_value(self.employee, 'casual', YEAR, MONTH), Decimal('1'))
        self.assertEqual(monthly_credit_value(self.employee, 'casual', YEAR, MONTH + 1), Decimal('2'))
        self.assertEqual(self.total(), Decimal('15'))

    def test_only_total_days_changes(self):
        self.set('3')
        self.balance.refresh_from_db()
        self.assertEqual((self.balance.used_days, self.balance.carried_forward), (Decimal('2'), Decimal('1')))
        self.assertEqual(self.casual.__class__.objects.get(pk=self.casual.pk).annual_days, Decimal('12'))
        other = LeaveBalance.objects.filter(employee=self.other_hyd).count()
        self.assertEqual(other, 0)

    def test_creates_current_year_balance_when_missing(self):
        result = self.set('1', employee=self.other_hyd)
        self.assertEqual(result['status'], 'created')
        self.assertEqual(LeaveBalance.objects.get(employee=self.other_hyd, leave_type='casual', year=YEAR).total_days, Decimal('1'))
        # zero for an employee with no balance creates nothing
        self.assertEqual(self.set('0', employee=self.emp_mum)['status'], 'unchanged')
        self.assertFalse(LeaveBalance.objects.filter(employee=self.emp_mum).exists())

    def test_validation(self):
        for bad in ('-1', '1.25', 'abc', '32'):
            with self.subTest(bad=bad), self.assertRaises(LeaveCreditError):
                self.set(bad)
        for month in (0, 13, 'x'):
            with self.subTest(month=month), self.assertRaises(LeaveCreditError):
                self.set('1', month=month)
        with self.assertRaises(LeaveCreditError):
            self.set('1', leave_type='lwp')
        with self.assertRaises(LeaveCreditError):
            self.set('1', year=YEAR + 1)
        self.employee.is_active = False
        self.employee.save(update_fields=['is_active'])
        with self.assertRaises(LeaveCreditError):
            self.set('1')
        self.assertEqual(self.total(), Decimal('12'))

    def test_reduction_cannot_go_below_used(self):
        self.set('5')
        self.balance.used_days = Decimal('16')
        self.balance.save(update_fields=['used_days'])
        with self.assertRaises(LeaveCreditError):
            self.set('0')
        self.assertEqual(self.total(), Decimal('17'))

    def test_monthly_rows_cannot_be_reversed_from_leave_credits(self):
        txn = self.set('1')['transaction']
        with self.assertRaises(LeaveCreditError):
            reverse_transaction(employee=self.employee, txn_id=txn.id, reason='x', actor=self.hr_hyd, idempotency_key=_key())

    def test_yearly_summary_still_reconciles(self):
        self.set('2')
        row = next(r for r in yearly_summary(self.employee, YEAR) if r['leave_type'] == 'casual')
        self.assertEqual(row['opening_or_allocated'] + row['carried_forward'] + row['manual_adjustments'], row['total_credits'])

    def test_dashboard_balance_cache_dropped_after_commit(self):
        key = f'dashboard:employee:leave_balance:{self.employee.id}:{YEAR}'
        cache.set(key, {'stale': True}, 300)
        with self.captureOnCommitCallbacks(execute=True):
            self.set('1')
        self.assertIsNone(cache.get(key))


class MonthlyLeaveApiTests(_Fixtures, TestCase):
    def setUp(self):
        self.make_fixtures()
        self.client = APIClient()

    def _login(self, user):
        resp = self.client.post(reverse('login'), {'company_code': TEST_COMPANY_CODE, 'email': user.email,
                                                   'password': PASSWORD}, format='json')
        self.assertEqual(resp.status_code, 200, resp.data)

    def _import(self, upload, dry_run=False, **extra):
        data = {'file': upload, 'year': YEAR, 'month': MONTH, **extra}
        if dry_run:
            data['dry_run'] = 'true'
        return self.client.post(reverse('leave-monthly-import'), data, format='multipart')

    # ── templates ──
    def test_sample_templates(self):
        self._login(self.hr_hyd)
        url = reverse('leave-monthly-import-sample')
        resp = self.client.get(url, {'format': 'csv'})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.decode('utf-8-sig').startswith('Employee Code,Leave Type,Monthly Leave Credit'))
        resp = self.client.get(url, {'format': 'xlsx'})
        self.assertEqual(resp.status_code, 200)
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        self.assertEqual([c.value for c in ws[1]], ['Employee Code', 'Leave Type', 'Monthly Leave Credit'])
        self.assertGreater(ws.max_row, 1)

    # ── import ──
    def test_preview_writes_nothing_and_reports_rows(self):
        self._login(self.hr_hyd)
        upload = _csv([
            ['CREMP01', 'casual', '1.5'],       # new
            ['CREMP02', 'Casual', '1'],         # new (label match, no balance yet → created on import)
            ['NOPE01', 'casual', '1'],          # unknown employee
            ['CREMP03', 'casual', '1'],         # other branch → hidden as not found
            ['CREMP01', 'bogus', '1'],          # bad leave type
            ['CREMP02', 'earned', '-1'],        # negative
            ['CREMP01', 'casual', '2'],         # duplicate of row 2
        ])
        resp = self._import(upload, dry_run=True)
        self.assertEqual(resp.status_code, 200, resp.data)
        data = resp.data['data']
        self.assertEqual(data['counts'], {'new': 2, 'update': 0, 'unchanged': 0, 'duplicate': 1, 'error': 4})
        self.assertIn('was not found', data['rows'][3]['message'])
        self.assertEqual(data['annual_entitlements'], {'casual': 12.0, 'earned': 15.0})
        self.assertFalse(LeaveCreditTransaction.objects.exists())
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.total_days, Decimal('12'))

    def test_import_partial_failure_and_reimport_is_noop(self):
        self._login(self.hr_hyd)
        rows = [['CREMP01', 'casual', '1.5'], ['NOPE01', 'casual', '1'], ['CREMP02', 'casual', '2']]
        resp = self._import(_csv(rows))
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['summary'], {'created': 2, 'updated': 0, 'skipped': 0, 'failed': 1, 'total': 3})
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.total_days, Decimal('13.5'))

        resp = self._import(_csv(rows))   # same file again: nothing double-credited
        self.assertEqual(resp.data['data']['summary'], {'created': 0, 'updated': 0, 'skipped': 2, 'failed': 1, 'total': 3})
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.total_days, Decimal('13.5'))
        self.assertEqual(LeaveCreditTransaction.objects.count(), 2)

        resp = self._import(_csv([['CREMP01', 'casual', '0.5']]))   # changed value updates by the difference
        self.assertEqual(resp.data['data']['summary']['updated'], 1)
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.total_days, Decimal('12.5'))

    def test_xlsx_import(self):
        self._login(self.hr_hyd)
        resp = self._import(_xlsx([('CREMP01', 'Casual Leave', 2)]))
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['summary']['created'], 1)
        self.assertEqual(monthly_credit_value(self.employee, 'casual', YEAR, MONTH), Decimal('2'))

    def test_bad_files_and_inputs(self):
        self._login(self.hr_hyd)
        self.assertEqual(self._import(_csv([['CREMP01', 'casual', '1']], header=('Code', 'Type', 'Days'))).status_code, 400)
        self.assertEqual(self._import(SimpleUploadedFile('x.txt', b'a,b')).status_code, 400)
        self.assertEqual(self._import(_csv([['CREMP01', 'casual', '1']]), month=13).status_code, 400)
        self.assertEqual(self._import(_csv([['CREMP01', 'casual', '1']]), year=YEAR + 1).status_code, 400)

    def test_concurrent_import_for_same_month_is_refused(self):
        self._login(self.hr_hyd)
        cache.add(f'leave:monthly_import:{YEAR}:{MONTH}', 'someone', 60)
        resp = self._import(_csv([['CREMP01', 'casual', '1']]))
        self.assertEqual(resp.status_code, 409)
        self.assertFalse(LeaveCreditTransaction.objects.exists())

    # ── manual entry ──
    def test_manual_add_edit_and_lookup(self):
        self._login(self.hr_hyd)
        url = reverse('leave-monthly-credits')
        payload = {'employee_id': 'CREMP01', 'year': YEAR, 'month': MONTH, 'leave_type': 'casual', 'credit': '1'}
        resp = self.client.post(url, payload, format='json')
        self.assertEqual((resp.status_code, resp.data['data']['status']), (200, 'created'))
        resp = self.client.post(url, {**payload, 'credit': '2.5'}, format='json')
        self.assertEqual((resp.data['data']['status'], resp.data['data']['delta']), ('updated', 1.5))
        resp = self.client.post(url, {**payload, 'credit': '2.5'}, format='json')
        self.assertEqual(resp.data['data']['status'], 'unchanged')
        self.assertEqual(self.client.post(url, {**payload, 'credit': '-1'}, format='json').status_code, 400)

        resp = self.client.get(url, {'employee_id': 'CREMP01', 'year': YEAR, 'month': MONTH})
        casual = next(r for r in resp.data['data']['leave_types'] if r['leave_type'] == 'casual')
        self.assertEqual(casual['monthly_credit'], 2.5)
        self.assertEqual(casual['balance']['total'], 14.5)
        self.assertEqual(casual['annual_days'], 12.0)

    # ── access ──
    def test_permissions_and_scope(self):
        url = reverse('leave-monthly-credits')
        payload = {'employee_id': 'CREMP03', 'year': YEAR, 'month': MONTH, 'leave_type': 'casual', 'credit': '1'}

        self._login(self.hr_hyd)   # other branch → not found
        self.assertEqual(self.client.post(url, payload, format='json').status_code, 404)
        self.assertEqual(self.client.post(url, {**payload, 'employee_id': 'CRHRH01'}, format='json').status_code, 404)  # self

        self._login(self.manager)  # only direct reports
        self.assertEqual(self.client.post(url, {**payload, 'employee_id': 'CREMP02'}, format='json').status_code, 404)
        self.assertEqual(self.client.post(url, {**payload, 'employee_id': 'CREMP01'}, format='json').status_code, 200)

        self._login(self.other_hyd)  # no leave.approve
        self.assertEqual(self.client.post(url, payload, format='json').status_code, 403)
        self.assertEqual(self._import(_csv([['CREMP01', 'casual', '1']])).status_code, 403)
        self.assertEqual(self.client.get(reverse('leave-monthly-import-sample')).status_code, 403)

        self._login(self.admin)      # settings.edit → all branches
        self.assertEqual(self.client.post(url, payload, format='json').status_code, 200)
        self.assertEqual(LeaveBalance.objects.get(employee=self.emp_mum, leave_type='casual', year=YEAR).total_days, Decimal('1'))


class MonthlyCreditConcurrencyTests(_Fixtures, TransactionTestCase):
    """Real commits and separate connections: concurrent saves of the same month apply once."""

    def setUp(self):
        self.make_fixtures()

    def tearDown(self):
        for conn in connections.all():
            conn.close()

    def test_concurrent_same_value_applies_once(self):
        results = []

        def run():
            from apps.tenants.models import Client
            connection.set_tenant(Client.objects.get(company_code=TEST_COMPANY_CODE))
            try:
                results.append(set_monthly_credit(employee=self.employee, leave_type='casual', year=YEAR,
                                                  month=MONTH, value='2', actor=self.hr_hyd)['status'])
            except Exception as exc:  # noqa: BLE001 — asserted below
                results.append(exc)
            finally:
                connections.close_all()

        threads = [threading.Thread(target=run) for _ in range(5)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertFalse([r for r in results if isinstance(r, Exception)], results)
        self.assertEqual(sorted(results), ['created'] + ['unchanged'] * 4)
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.total_days, Decimal('14'))
        self.assertEqual(LeaveCreditTransaction.objects.count(), 1)
