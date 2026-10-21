"""
Bulk asset import (CSV / Excel) — Settings → Assets → Bulk Import.
Real authenticated requests through the URL/view/permission stack.
"""
from __future__ import annotations

import base64
import csv
import datetime
import io

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.factories import make_role, make_user
from apps.accounts.models import AuditLog
from apps.assets.models import Asset, AssetCategory, AssetType
from apps.assets.services_import import HEADERS
from apps.assets.tests import _login, _make_branch


def _csv(rows, headers=None, name='assets.csv'):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers or HEADERS)
    for r in rows:
        w.writerow([r.get(h, '') for h in (headers or HEADERS)])
    return SimpleUploadedFile(name, buf.getvalue().encode('utf-8'), content_type='text/csv')


def _xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h, '') for h in HEADERS])
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile('assets.xlsx', out.getvalue(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class AssetBulkImportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.branch_a = _make_branch('ImpBranchA')
        self.branch_b = _make_branch('ImpBranchB')
        self.category = AssetCategory.objects.create(name='Import IT')
        AssetType.objects.create(name='Import Laptop', category=self.category)
        admin_role = make_role('imp_admin', permission_codenames=['assets.view', 'assets.create', 'settings.edit'])
        self.admin = make_user('imp.admin@test.com', role=admin_role, employee_id='IMPADM1',
                               full_name='Import Admin', branch=self.branch_a.branch_name)
        branch_role = make_role('imp_branch', permission_codenames=['assets.view', 'assets.create'])
        self.branch_user = make_user('imp.branch@test.com', role=branch_role, employee_id='IMPBR1',
                                     full_name='Branch User', branch=self.branch_a.branch_name)
        viewer_role = make_role('imp_viewer', permission_codenames=['assets.view'])
        self.viewer = make_user('imp.viewer@test.com', role=viewer_role, employee_id='IMPVW1',
                                full_name='Viewer', branch=self.branch_a.branch_name)
        self.client = APIClient()

    def row(self, tag, **over):
        base = {'Asset Tag': tag, 'Asset Name': f'Laptop {tag}', 'Category': 'Import IT',
                'Asset Type': 'Import Laptop', 'Branch': self.branch_a.branch_code}
        base.update(over)
        return base

    def _import(self, upload, dry_run=False):
        data = {'file': upload}
        if dry_run:
            data['dry_run'] = 'true'
        return self.client.post(reverse('asset-import'), data, format='multipart')

    # 1 + 2 ── successful CSV and Excel imports
    def test_csv_preview_then_import(self):
        _login(self.client, self.admin.email)
        rows = [
            self.row('IMP-001', **{'Brand': 'Dell', 'Model': 'Latitude', 'Serial Number': 'SN-1',
                                   'Purchase Date': '15-01-2026', 'Purchase Price': '75,000.50',
                                   'Warranty Start Date': '2026-01-15', 'Warranty End Date': '2029-01-14',
                                   'Condition': 'Good', 'Vendor': 'V1', 'Description': 'First'}),
            self.row('IMP-002', **{'Branch': self.branch_b.branch_name}),       # by name, another branch (admin)
        ]
        resp = self._import(_csv(rows), dry_run=True)
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['counts'], {'ready': 2, 'duplicate': 0, 'error': 0})
        self.assertFalse(Asset.objects.filter(asset_tag__startswith='IMP-').exists())

        resp = self._import(_csv(rows))
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['summary'], {'total': 2, 'created': 2, 'duplicate': 0, 'failed': 0, 'skipped': 0})
        a = Asset.objects.get(asset_tag='IMP-001')
        self.assertEqual((a.purchase_date, str(a.purchase_price), a.condition, a.status, a.serial_number, a.created_by_id),
                         (datetime.date(2026, 1, 15), '75000.50', 'good', 'available', 'SN-1', self.admin.id))
        self.assertEqual(Asset.objects.get(asset_tag='IMP-002').branch, self.branch_b)
        self.assertEqual(Asset.objects.get(asset_tag='IMP-002').condition, 'new')   # blank -> form default
        self.assertEqual(AuditLog.objects.filter(action='asset_created', changes__source='bulk_import').count(), 2)
        self.assertTrue(AuditLog.objects.filter(action='asset_bulk_import').exists())

    def test_xlsx_import_with_real_date_cells(self):
        _login(self.client, self.admin.email)
        rows = [self.row('IMP-X1', **{'Purchase Date': datetime.date(2026, 3, 1), 'Purchase Price': 1200})]
        resp = self._import(_xlsx(rows))
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['summary']['created'], 1)
        a = Asset.objects.get(asset_tag='IMP-X1')
        self.assertEqual((a.purchase_date, str(a.purchase_price)), (datetime.date(2026, 3, 1), '1200.00'))

    # 3 ── templates
    def test_sample_templates_download_and_validate(self):
        import openpyxl
        _login(self.client, self.admin.email)
        resp = self.client.get(reverse('asset-import-sample'), {'format': 'csv'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment; filename="asset_import_sample.csv"', resp['Content-Disposition'])
        lines = list(csv.reader(io.StringIO(resp.content.decode('utf-8-sig'))))
        self.assertEqual(lines[0], HEADERS)
        self.assertEqual(len(lines), 3)
        resp = self.client.get(reverse('asset-import-sample'), {'format': 'xlsx'})
        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual([c.value for c in wb['Assets'][1]], HEADERS)
        self.assertIn('Instructions', wb.sheetnames)
        # the template's own example rows pass validation as-is
        resp = self._import(SimpleUploadedFile('sample.xlsx', resp.content), dry_run=True)
        self.assertEqual(resp.data['data']['counts']['ready'], 2, resp.data)

    # 4 ── headers
    def test_missing_or_unknown_headers_rejected(self):
        _login(self.client, self.admin.email)
        resp = self._import(_csv([{'Asset Tag': 'X'}], headers=['Asset Tag', 'Asset Name']))
        self.assertEqual(resp.status_code, 400)
        self.assertIn('Missing required column(s): Category, Asset Type, Branch', resp.data['message'])
        resp = self._import(_csv([{}], headers=HEADERS + ['Serail Number']))
        self.assertEqual(resp.status_code, 400)
        self.assertIn('Unknown column(s): Serail Number', resp.data['message'])
        self.assertEqual(self._import(SimpleUploadedFile('a.txt', b'x')).status_code, 400)

    # 5 + 7 ── required fields, formats, references
    def test_row_level_errors_name_the_column(self):
        _login(self.client, self.admin.email)
        rows = [
            self.row('', **{'Asset Name': ''}),
            self.row('IMP-E2', **{'Purchase Date': '31/31/2026', 'Purchase Price': '-5', 'Condition': 'broken'}),
            self.row('IMP-E3', **{'Category': 'Nope'}),
            self.row('IMP-E4', **{'Asset Type': 'Toaster'}),
            self.row('IMP-E5', **{'Branch': 'NOWHERE'}),
            self.row('IMP-E6', **{'Warranty Start Date': '2026-05-01', 'Warranty End Date': '2026-04-01'}),
            self.row('IMP-E7', **{'Category': 'Other', 'Category Other': '', 'Asset Type': 'Other'}),
        ]
        data = self._import(_csv(rows), dry_run=True).data['data']
        self.assertEqual(data['counts']['error'], 7)
        cols = {r['row']: {e['column'] for e in r['errors']} for r in data['rows']}
        self.assertTrue({'Asset Tag', 'Asset Name'} <= cols[2])
        self.assertEqual(cols[3], {'Purchase Date', 'Purchase Price', 'Condition'})
        self.assertEqual(cols[4], {'Category'})
        self.assertEqual(cols[5], {'Asset Type'})
        self.assertEqual(cols[6], {'Branch'})
        self.assertEqual(cols[7], {'Warranty End Date'})
        self.assertEqual(cols[8], {'Category Other'})
        self.assertEqual(self._import(_csv(rows)).data['data']['summary']['created'], 0)

    # 6 + 8 + 10 ── duplicates, partial failure, retry safety
    def test_duplicates_partial_failure_and_retry(self):
        Asset.objects.create(asset_tag='IMP-OLD', asset_name='Old', category='Import IT', asset_type='Import Laptop',
                             branch=self.branch_a, serial_number='SN-OLD')
        _login(self.client, self.admin.email)
        rows = [
            self.row('IMP-D1', **{'Serial Number': 'SN-D1'}),
            self.row('imp-old'),                                          # already in the system (case-insensitive)
            self.row('IMP-D1'),                                           # repeated in the file
            self.row('IMP-D2', **{'Serial Number': 'SN-OLD'}),            # serial already in the system
            self.row('IMP-D3', **{'Category': 'Nope'}),                   # invalid
        ]
        resp = self._import(_csv(rows))
        s = resp.data['data']['summary']
        self.assertEqual(s, {'total': 5, 'created': 1, 'duplicate': 3, 'failed': 1, 'skipped': 0})
        report = base64.b64decode(resp.data['data']['error_report_csv']).decode('utf-8-sig')
        lines = list(csv.reader(io.StringIO(report)))
        self.assertEqual(lines[0], ['Row', 'Asset Tag', 'Status', 'Column', 'Reason'])
        self.assertEqual(len(lines) - 1, 4)
        self.assertNotIn('Traceback', report)

        before = Asset.objects.count()
        s = self._import(_csv(rows)).data['data']['summary']                # retried: nothing created twice
        self.assertEqual((s['created'], s['duplicate']), (0, 4))
        self.assertEqual(Asset.objects.count(), before)

    # 9 ── permissions and branch scope
    def test_permissions_and_branch_scope(self):
        _login(self.client, self.viewer.email)
        self.assertEqual(self._import(_csv([self.row('IMP-P0')])).status_code, 403)
        self.assertEqual(self.client.get(reverse('asset-import-sample')).status_code, 403)

        _login(self.client, self.branch_user.email)
        rows = [self.row('IMP-P1'), self.row('IMP-P2', **{'Branch': self.branch_b.branch_code})]
        data = self._import(_csv(rows)).data['data']
        self.assertEqual(data['summary']['created'], 1)
        other = next(r for r in data['rows'] if r['asset_tag'] == 'IMP-P2')
        self.assertEqual(other['errors'][0], {'column': 'Branch', 'message': 'You can only import assets for your own branch.'})
        self.assertFalse(Asset.objects.filter(asset_tag='IMP-P2').exists())
        sample = self.client.get(reverse('asset-import-sample'), {'format': 'csv'}).content.decode('utf-8-sig')
        self.assertIn(self.branch_a.branch_code, sample)                  # template uses the user's own branch

    # 10 ── the single Add Asset flow is unchanged
    def test_single_add_asset_still_works(self):
        _login(self.client, self.admin.email)
        resp = self.client.post(reverse('asset-list-create'), {
            'asset_tag': 'IMP-SINGLE', 'asset_name': 'Single', 'category': 'Import IT', 'asset_type': 'Import Laptop',
            'branch': self.branch_a.pk,
        }, format='json')
        self.assertEqual(resp.status_code, 201, resp.data)
