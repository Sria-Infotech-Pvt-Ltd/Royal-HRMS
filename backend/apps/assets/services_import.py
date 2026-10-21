"""
Bulk asset import (CSV / Excel) for Settings → Assets.

Every row is validated by the existing AssetSerializer — the same rules as
the single Add Asset form (required fields, unique Asset Tag / Serial
Number, active Category / Asset Type master data, "Other" handling,
warranty dates) — plus the same branch rule as AssetListCreateView.post().
The preview (dry run) and the confirmed import run the same validation;
confirmed rows are saved one transaction each, so a failed row never
affects the others. Re-importing a file creates nothing twice: an Asset Tag
that already exists is reported as a duplicate and skipped.

Columns are exactly the Add Asset form's fields. There is no Status column:
new assets start Available, as in the form; Assigned / Under Repair are set
only through the assignment and maintenance workflows.
"""
from __future__ import annotations

import base64
import csv
import datetime
import io
from decimal import Decimal, InvalidOperation

from django.db import IntegrityError, transaction

from apps.assets.models import Asset, AssetCategory, AssetType

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000

# (header, serializer field, required, notes)
COLUMNS = [
    ('Asset Tag',           'asset_tag',           True,  'Unique, e.g. AST-0101'),
    ('Asset Name',          'asset_name',          True,  ''),
    ('Category',            'category',            True,  'An active Asset Category, or "Other"'),
    ('Category Other',      'category_other',      False, 'Required when Category is "Other"'),
    ('Asset Type',          'asset_type',          True,  'An active Asset Type of that Category, or "Other"'),
    ('Asset Type Other',    'asset_type_other',    False, 'Required when Asset Type is "Other"'),
    ('Branch',              'branch',              True,  'Branch code or branch name'),
    ('Brand',               'brand',               False, ''),
    ('Model',               'model',               False, ''),
    ('Serial Number',       'serial_number',       False, 'Unique when given'),
    ('Purchase Date',       'purchase_date',       False, 'YYYY-MM-DD or DD-MM-YYYY'),
    ('Purchase Price',      'purchase_price',      False, 'Number, e.g. 75000 or 75000.50'),
    ('Vendor',              'vendor',              False, ''),
    ('Warranty Start Date', 'warranty_start_date', False, 'YYYY-MM-DD or DD-MM-YYYY'),
    ('Warranty End Date',   'warranty_end_date',   False, 'Not before the start date'),
    ('Condition',           'condition',           False, 'new, good, fair or damaged (blank = new, as in the form)'),
    ('Description',         'description',         False, ''),
]
HEADERS = [c[0] for c in COLUMNS]
_FIELD_BY_HEADER = {c[0].lower(): c[1] for c in COLUMNS}
_HEADER_BY_FIELD = {c[1]: c[0] for c in COLUMNS}
_REQUIRED = [c[0] for c in COLUMNS if c[2]]
_DATE_FIELDS = ('purchase_date', 'warranty_start_date', 'warranty_end_date')
_DATE_FORMATS = ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d')
_CONDITIONS = {value: value for value, _ in Asset.CONDITION_CHOICES}
_CONDITIONS.update({label.lower(): value for value, label in Asset.CONDITION_CHOICES})


class ImportFileError(Exception):
    """The file as a whole can't be imported (type, size, headers, row count)."""


# ─── Parsing ──────────────────────────────────────────────────────────────────

def _cell_text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(uploaded) -> list[list[str]]:
    name = (uploaded.name or '').lower()
    if uploaded.size > MAX_FILE_BYTES:
        raise ImportFileError('File must not exceed 5 MB.')
    if name.endswith('.csv'):
        try:
            text = uploaded.read().decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ImportFileError('The CSV file must be UTF-8 encoded.')
        return [[(c or '').strip() for c in row] for row in csv.reader(io.StringIO(text))]
    if name.endswith('.xlsx'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(uploaded, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            wb.close()
            return rows
        except Exception:
            raise ImportFileError('Could not read the Excel file. Save it as .xlsx and try again.')
    raise ImportFileError('Unsupported file type. Upload a .csv or .xlsx file.')


def parse_file(uploaded) -> list[dict]:
    """[{'row': n, <field>: text, ...}] with n = spreadsheet row number. Raises ImportFileError."""
    raw = _read_rows(uploaded)
    raw = [r for r in raw if any(c for c in r)]
    if not raw:
        raise ImportFileError('The file is empty.')
    headers = [h.strip() for h in raw[0]]
    unknown = [h for h in headers if h and h.lower() not in _FIELD_BY_HEADER]
    if unknown:
        raise ImportFileError(f'Unknown column(s): {", ".join(unknown)}. Use the sample template headers.')
    present = {h.lower() for h in headers if h}
    missing = [h for h in _REQUIRED if h.lower() not in present]
    if missing:
        raise ImportFileError(f'Missing required column(s): {", ".join(missing)}.')
    body = raw[1:]
    if not body:
        raise ImportFileError('The file has no data rows.')
    if len(body) > MAX_ROWS:
        raise ImportFileError(f'The file has {len(body)} rows; import at most {MAX_ROWS} at a time.')
    rows = []
    for offset, values in enumerate(body, start=2):
        row = {'row': offset}
        for i, header in enumerate(headers):
            if header:
                row[_FIELD_BY_HEADER[header.lower()]] = values[i] if i < len(values) else ''
        rows.append(row)
    return rows


# ─── Validation ───────────────────────────────────────────────────────────────

def _branch_lookup():
    from apps.branch.models import Branch
    lookup = {}
    for b in Branch.objects.filter(status=Branch.STATUS_ACTIVE):
        lookup[b.branch_code.strip().lower()] = b
        lookup.setdefault(b.branch_name.strip().lower(), b)
    return lookup


def _parse_date(text: str):
    for fmt in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _first_errors(errors: dict) -> list[dict]:
    out = []
    for field, messages in errors.items():
        message = messages[0] if isinstance(messages, list) and messages else messages
        out.append({'column': _HEADER_BY_FIELD.get(field, field), 'message': str(message)})
    return out


def plan_import(rows: list[dict], *, user, is_admin: bool, user_branch) -> dict:
    """Validate every row without saving. Status per row: ready / duplicate / error."""
    from apps.assets.serializers import AssetSerializer

    branches = _branch_lookup()
    seen_tags, seen_serials = set(), set()
    planned = []
    for row in rows:
        out = {'row': row['row'], 'asset_tag': row.get('asset_tag', ''), 'asset_name': row.get('asset_name', ''),
               'branch': row.get('branch', ''), 'status': 'error', 'errors': [], '_data': None}
        planned.append(out)
        errors = out['errors']

        data = {field: (row.get(field) or '').strip() for field in _HEADER_BY_FIELD if field != 'branch'}
        for field in _DATE_FIELDS:
            if data[field]:
                parsed = _parse_date(data[field])
                if parsed is None:
                    errors.append({'column': _HEADER_BY_FIELD[field], 'message': f'"{data[field]}" is not a valid date (use YYYY-MM-DD).'})
                data[field] = parsed.isoformat() if parsed else None
            else:
                data[field] = None
        if data['purchase_price']:
            try:
                price = Decimal(data['purchase_price'].replace(',', ''))
                if not price.is_finite() or price < 0:
                    raise InvalidOperation
                data['purchase_price'] = str(price)
            except (InvalidOperation, ValueError):
                errors.append({'column': 'Purchase Price', 'message': f'"{data["purchase_price"]}" is not a valid non-negative number.'})
                data['purchase_price'] = None
        else:
            data['purchase_price'] = None
        condition = data['condition'].lower()
        if condition and condition not in _CONDITIONS:
            errors.append({'column': 'Condition', 'message': f'"{data["condition"]}" is not one of new, good, fair, damaged.'})
        data['condition'] = _CONDITIONS.get(condition) if condition else Asset.CONDITION_NEW   # blank = form default

        raw_branch = (row.get('branch') or '').strip()
        branch = branches.get(raw_branch.lower()) if raw_branch else None
        if not raw_branch:
            errors.append({'column': 'Branch', 'message': 'Branch is required.'})
        elif branch is None:
            errors.append({'column': 'Branch', 'message': f'"{raw_branch}" is not an active branch code or name.'})
        elif not is_admin and (user_branch is None or branch.pk != user_branch.pk):
            errors.append({'column': 'Branch', 'message': 'You can only import assets for your own branch.'})
        if branch is not None:
            out['branch'] = branch.branch_name
            data['branch'] = branch.pk

        if errors:
            continue
        # Only rows that would otherwise be imported claim a tag/serial in this file.
        tag_key = data['asset_tag'].lower()
        serial_key = data['serial_number'].lower()
        dup_tag = bool(tag_key) and tag_key in seen_tags
        dup_serial = bool(serial_key) and serial_key in seen_serials
        if not (dup_tag or dup_serial):
            if tag_key:
                seen_tags.add(tag_key)
            if serial_key:
                seen_serials.add(serial_key)
        if dup_tag or dup_serial:
            out['status'] = 'duplicate'
            out['errors'].append({'column': 'Asset Tag' if dup_tag else 'Serial Number',
                                  'message': 'Repeated in this file; only the first row is imported.'})
            continue
        serializer = AssetSerializer(data=data)
        if serializer.is_valid():
            out['status'] = 'ready'
            out['_data'] = data
            continue
        row_errors = _first_errors(serializer.errors)
        already = [e for e in row_errors if e['column'] in ('Asset Tag', 'Serial Number') and 'already exists' in e['message']]
        if already and len(already) == len(row_errors):
            out['status'] = 'duplicate'            # already in the system (e.g. a retried import)
        errors.extend(row_errors)

    counts = {s: sum(1 for p in planned if p['status'] == s) for s in ('ready', 'duplicate', 'error')}
    return {'rows': planned, 'counts': counts}


def apply_import(plan: dict, *, user, ip_address=None) -> dict:
    """Save every 'ready' row through AssetSerializer, one transaction each."""
    from apps.accounts.models import AuditLog
    from apps.assets.serializers import AssetSerializer

    created = 0
    for row in plan['rows']:
        data = row.pop('_data', None)
        if row['status'] != 'ready':
            continue
        serializer = AssetSerializer(data=data)
        if not serializer.is_valid():                      # changed since the preview
            row['errors'] = _first_errors(serializer.errors)
            row['status'] = 'duplicate' if all('already exists' in e['message'] for e in row['errors']) else 'error'
            continue
        try:
            with transaction.atomic():
                asset = serializer.save(created_by=user)
                AuditLog.objects.create(
                    user=user, action='asset_created', module='assets', object_id=str(asset.pk),
                    changes={'asset_tag': asset.asset_tag, 'asset_name': asset.asset_name, 'source': 'bulk_import'},
                    branch=asset.branch.branch_name, ip_address=ip_address,
                )
        except IntegrityError:                             # created concurrently
            row['status'] = 'duplicate'
            row['errors'] = [{'column': 'Asset Tag', 'message': f'Asset Tag "{row["asset_tag"]}" already exists.'}]
            continue
        row['status'] = 'created'
        created += 1
    for row in plan['rows']:
        row.pop('_data', None)
    summary = {
        'total':      len(plan['rows']),
        'created':    created,
        'duplicate':  sum(1 for r in plan['rows'] if r['status'] == 'duplicate'),
        'failed':     sum(1 for r in plan['rows'] if r['status'] == 'error'),
    }
    summary['skipped'] = summary['total'] - summary['created'] - summary['duplicate'] - summary['failed']
    return {'rows': plan['rows'], 'summary': summary, 'error_report_csv': error_report_csv(plan['rows'])}


def strip_internal(plan: dict) -> dict:
    for row in plan['rows']:
        row.pop('_data', None)
    plan['error_report_csv'] = error_report_csv(plan['rows'])
    return plan


def error_report_csv(rows: list[dict]) -> str | None:
    """Base64 CSV of rows that weren't imported (row, asset tag, column, reason) — no internal details."""
    problems = [r for r in rows if r['status'] in ('error', 'duplicate')]
    if not problems:
        return None
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(['Row', 'Asset Tag', 'Status', 'Column', 'Reason'])
    for r in problems:
        for e in r['errors'] or [{'column': '', 'message': ''}]:
            reason = e['message']
            if reason[:1] in ('=', '+', '-', '@'):
                reason = "'" + reason
            writer.writerow([r['row'], r['asset_tag'], r['status'], e['column'], reason])
    return base64.b64encode(buf.getvalue().encode('utf-8-sig')).decode('ascii')


# ─── Sample templates ─────────────────────────────────────────────────────────

def sample_rows(*, branch_code: str) -> list[list[str]]:
    """Two clearly-labelled example rows using a real active category/type and branch."""
    category = AssetCategory.objects.filter(is_active=True).order_by('name').first()
    asset_type = (AssetType.objects.filter(is_active=True, category=category).order_by('name').first()
                  if category else None)
    cat = category.name if category else 'Other'
    typ = asset_type.name if asset_type else 'Other'
    cat_other = '' if category else 'Example category'
    typ_other = '' if asset_type else 'Example type'
    example = {
        'asset_tag': 'EXAMPLE-001', 'asset_name': 'Example asset (replace or delete this row)',
        'category': cat, 'category_other': cat_other, 'asset_type': typ, 'asset_type_other': typ_other,
        'branch': branch_code, 'brand': 'Dell', 'model': 'Latitude 5440', 'serial_number': 'EXAMPLE-SN-001',
        'purchase_date': '2026-01-15', 'purchase_price': '75000', 'vendor': 'Example Vendor',
        'warranty_start_date': '2026-01-15', 'warranty_end_date': '2029-01-14', 'condition': 'new',
        'description': 'Example row: replace with real data or delete before importing',
    }
    minimal = {k: '' for k in example}
    minimal.update({'asset_tag': 'EXAMPLE-002', 'asset_name': 'Example asset, required columns only',
                    'category': cat, 'category_other': cat_other, 'asset_type': typ,
                    'asset_type_other': typ_other, 'branch': branch_code})
    return [[r[c[1]] for c in COLUMNS] for r in (example, minimal)]


def sample_xlsx(rows: list[list[str]]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    from core.file_utils import build_sample_xlsx
    wb = openpyxl.load_workbook(io.BytesIO(build_sample_xlsx(HEADERS, rows, 'Assets')))
    info = wb.create_sheet('Instructions')
    info.append(['Column', 'Required', 'Notes'])
    for cell in info[1]:
        cell.font = Font(bold=True)
    for header, _field, required, notes in COLUMNS:
        info.append([header, 'Yes' if required else 'No', notes])
    info.append([])
    info.append(['Keep the header row exactly as in the Assets sheet. One asset per row; up to 1000 rows per file.'])
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 10
    info.column_dimensions['C'].width = 60
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
