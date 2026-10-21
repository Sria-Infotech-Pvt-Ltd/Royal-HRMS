"""
Bulk asset assignment (CSV / Excel) — Settings → Assets → Bulk Assign — and
the shared assign_asset service it uses with single Assign Asset.
Real authenticated requests through the URL/view/permission stack.
"""
from __future__ import annotations

import base64
import csv
import datetime
import io
from unittest import mock

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.factories import make_role, make_user
from apps.accounts.models import AuditLog, User
from apps.assets.models import Asset, AssetAssignment
from apps.assets.services_assign_import import HEADERS, apply_assignments, parse_file, plan_assignments
from apps.assets.services_assignment import assign_asset
from apps.assets.tests import _login, _make_branch


def _csv(rows, headers=None, name='assign.csv'):
    headers = headers or HEADERS
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow([r.get(h, '') for h in headers])
    return SimpleUploadedFile(name, buf.getvalue().encode('utf-8'), content_type='text/csv')


def _xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h, '') for h in HEADERS])
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile('assign.xlsx', out.getvalue())


def _row(tag, emp, **over):
    base = {'Asset Tag': tag, 'Employee ID': emp, 'Assigned Date': '2026-02-01', 'Condition at Assignment': 'good'}
    base.update(over)
    return base


class AssetBulkAssignTests(TestCase):
    def setUp(self):
        cache.clear()
        self.branch_a = _make_branch('AsgBranchA')
        self.branch_b = _make_branch('AsgBranchB')
        a, b = self.branch_a.branch_name, self.branch_b.branch_name
        complete = User.ONBOARDING_COMPLETE

        admin_role = make_role('asg_admin', permission_codenames=['assets.view', 'assets.edit', 'settings.edit'])
        self.admin = make_user('asg.admin@test.com', role=admin_role, employee_id='ASGADM', full_name='Asg Admin',
                               branch=a, onboarding_status=complete)
        branch_role = make_role('asg_branch', permission_codenames=['assets.view', 'assets.edit'])
        self.branch_user = make_user('asg.branch@test.com', role=branch_role, employee_id='ASGBR', full_name='Asg Branch',
                                     branch=a, onboarding_status=complete)
        manager_role = make_role('asg_manager', permission_codenames=['assets.view', 'assets.edit'], can_manage_team=True)
        self.manager = make_user('asg.mgr@test.com', role=manager_role, employee_id='ASGMGR', full_name='Asg Manager',
                                 branch=a, onboarding_status=complete)
        creator_role = make_role('asg_creator', permission_codenames=['assets.view', 'assets.create'])
        self.creator = make_user('asg.creator@test.com', role=creator_role, employee_id='ASGCR', full_name='Asg Creator',
                                 branch=a, onboarding_status=complete)

        emp_role = make_role('asg_emp')
        mk = lambda email, code, branch, **kw: make_user(email, role=emp_role, employee_id=code, branch=branch,
                                                          full_name=code, **{'onboarding_status': complete, **kw})
        self.emp_a1 = mk('a1@test.com', 'EMPA1', a, reporting_manager=self.manager)
        self.emp_a2 = mk('a2@test.com', 'EMPA2', a)
        self.emp_b1 = mk('b1@test.com', 'EMPB1', b)
        self.inactive = mk('in@test.com', 'EMPOFF', a, is_active=False)
        self.pending = mk('pend@test.com', 'EMPNEW', a, onboarding_status=User.ONBOARDING_PENDING)
        mk('dup1@test.com', 'DUPX', a)
        mk('dup2@test.com', 'dupx', a)

        def asset(tag, branch, **kw):
            return Asset.objects.create(asset_tag=tag, asset_name=f'Laptop {tag}', category='IT', asset_type='Laptop',
                                        branch=branch, **kw)
        self.a1, self.a2, self.a3, self.a4 = (asset(f'ASG-A{i}', self.branch_a) for i in range(1, 5))
        self.b1 = asset('ASG-B1', self.branch_b)
        self.repair = asset('ASG-REP', self.branch_a, status=Asset.STATUS_UNDER_REPAIR)
        self.client = APIClient()

    def _import(self, upload, dry_run=False):
        data = {'file': upload}
        if dry_run:
            data['dry_run'] = 'true'
        return self.client.post(reverse('asset-assign-import'), data, format='multipart')

    def _errors(self, data):
        return {r['row']: [(e['column'], e['message']) for e in r['errors']] for r in data['rows']}

    # ── valid CSV / Excel assignments ─────────────────────────────────────
    def test_csv_preview_then_confirm(self):
        _login(self.client, self.admin.email)
        rows = [
            _row('asg-a1', 'empa1'),                                         # both matched case-insensitively
            _row('ASG-A2', 'EMPA2', **{'Assigned Date': '05-02-2026', 'Condition at Assignment': 'New',
                                       'Expected Return Date': '2027-02-04', 'Remarks': 'Joining kit'}),
            _row('ASG-B1', 'EMPB1'),                                         # admin: another branch is fine
        ]
        resp = self._import(_csv(rows), dry_run=True)
        self.assertEqual(resp.status_code, 200, resp.data)
        data = resp.data['data']
        self.assertEqual(data['counts'], {'ready': 3, 'already_assigned': 0, 'error': 0})
        self.assertEqual((data['rows'][0]['asset_tag'], data['rows'][0]['employee_id']), ('ASG-A1', 'EMPA1'))
        self.assertFalse(AssetAssignment.objects.exists())                  # preview saves nothing

        resp = self._import(_csv(rows))
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['summary'],
                         {'total': 3, 'assigned': 3, 'already_assigned': 0, 'failed': 0, 'skipped': 0})
        self.assertIsNone(resp.data['data']['error_report_csv'])
        asg = AssetAssignment.objects.get(asset=self.a2)
        self.assertEqual((asg.employee, asg.assigned_date, asg.condition_at_assignment, asg.expected_return_date,
                          asg.assign_remarks, asg.assigned_by, asg.status),
                         (self.emp_a2, datetime.date(2026, 2, 5), 'new', datetime.date(2027, 2, 4),
                          'Joining kit', self.admin, AssetAssignment.STATUS_ASSIGNED))
        self.a1.refresh_from_db()
        self.assertEqual((self.a1.status, self.a1.condition), (Asset.STATUS_ASSIGNED, 'good'))
        self.assertEqual(AssetAssignment.objects.get(asset=self.b1).employee, self.emp_b1)
        self.assertEqual(AuditLog.objects.filter(action='asset_assigned', changes__source='bulk_import').count(), 3)
        self.assertTrue(AuditLog.objects.filter(action='asset_bulk_assign').exists())

    def test_xlsx_with_real_date_cells(self):
        _login(self.client, self.admin.email)
        resp = self._import(_xlsx([_row('ASG-A1', 'EMPA1', **{'Assigned Date': datetime.date(2026, 3, 1),
                                                              'Expected Return Date': datetime.date(2026, 9, 1)})]))
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['data']['summary']['assigned'], 1)
        asg = AssetAssignment.objects.get(asset=self.a1)
        self.assertEqual((asg.assigned_date, asg.expected_return_date), (datetime.date(2026, 3, 1), datetime.date(2026, 9, 1)))

    # ── templates ─────────────────────────────────────────────────────────
    def test_sample_templates(self):
        import openpyxl
        _login(self.client, self.admin.email)
        resp = self.client.get(reverse('asset-assign-import-sample'), {'format': 'csv'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('filename="asset_assignment_sample.csv"', resp['Content-Disposition'])
        lines = list(csv.reader(io.StringIO(resp.content.decode('utf-8-sig'))))
        self.assertEqual(lines[0], ['Asset Tag', 'Employee ID', 'Assigned Date', 'Condition at Assignment',
                                    'Expected Return Date', 'Remarks'])
        self.assertEqual(len(lines), 3)
        self.assertTrue(all(line[0].startswith('EXAMPLE-') for line in lines[1:]))

        resp = self.client.get(reverse('asset-assign-import-sample'), {'format': 'xlsx'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('filename="asset_assignment_sample.xlsx"', resp['Content-Disposition'])
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual([c.value for c in wb['Assignments'][1]], HEADERS)
        self.assertIn('Instructions', wb.sheetnames)
        # the template parses, and its placeholder rows fail only on the asset/employee lookup — never assign
        data = self._import(SimpleUploadedFile('sample.xlsx', resp.content), dry_run=True).data['data']
        self.assertEqual(data['counts']['error'], 2)
        for row in data['rows']:
            self.assertEqual({e['column'] for e in row['errors']}, {'Asset Tag', 'Employee ID'})
        self.assertFalse(AssetAssignment.objects.exists())

    # ── file-level checks ────────────────────────────────────────────────
    def test_missing_or_unknown_headers_and_bad_file(self):
        _login(self.client, self.admin.email)
        resp = self._import(_csv([{'Asset Tag': 'X'}], headers=['Asset Tag', 'Remarks']))
        self.assertEqual(resp.status_code, 400)
        self.assertIn('Missing required column(s): Employee ID, Assigned Date, Condition at Assignment',
                      resp.data['message'])
        resp = self._import(_csv([{}], headers=HEADERS + ['Employee Name']))
        self.assertEqual(resp.status_code, 400)
        self.assertIn('Unknown column(s): Employee Name', resp.data['message'])
        self.assertEqual(self._import(SimpleUploadedFile('a.pdf', b'x')).status_code, 400)
        self.assertEqual(self._import(_csv([])).status_code, 400)                # header only

    # ── row validation ────────────────────────────────────────────────────
    def test_row_validation(self):
        _login(self.client, self.admin.email)
        rows = [
            _row('', ''),                                                                  # 2
            _row('ASG-A1', 'EMPA1', **{'Assigned Date': '', 'Condition at Assignment': ''}),  # 3
            _row('ASG-A1', 'EMPA1', **{'Assigned Date': '31/31/2026', 'Condition at Assignment': 'broken'}),  # 4
            _row('ASG-A2', 'EMPA1', **{'Expected Return Date': '2026-01-01'}),             # 5
            _row('NOPE-1', 'EMPA1'),                                                       # 6
            _row('ASG-A3', 'NOBODY'),                                                      # 7
            _row('ASG-A3', 'EMPOFF'),                                                      # 8 (repeats row 7's tag)
            _row('ASG-A4', 'EMPOFF'),                                                      # 9
            _row('ASG-B1', 'EMPNEW'),                                                      # 10
            _row('ASG-REP', 'DUPX'),                                                       # 11
        ]
        data = self._import(_csv(rows), dry_run=True).data['data']
        self.assertEqual(data['counts'], {'ready': 0, 'already_assigned': 0, 'error': 10})
        errs = self._errors(data)
        self.assertEqual(errs[2], [('Asset Tag', 'Asset Tag is required.'), ('Employee ID', 'Employee ID is required.')])
        self.assertIn(('Assigned Date', 'Assigned Date is required.'), errs[3])
        self.assertIn(('Condition at Assignment', 'Condition at Assignment is required.'), errs[3])
        repeated = ('Asset Tag', 'This Asset Tag appears in more than one row; list each asset once.')
        self.assertIn(repeated, errs[3])
        self.assertIn(repeated, errs[4])
        self.assertIn(('Assigned Date', '"31/31/2026" is not a valid date (use YYYY-MM-DD).'), errs[4])
        self.assertIn(('Condition at Assignment', '"broken" is not one of new, good, fair, damaged.'), errs[4])
        self.assertEqual(errs[5], [('Expected Return Date', 'Expected return date cannot be before the assigned date.')])
        self.assertEqual(errs[6], [('Asset Tag', 'No asset with Asset Tag "NOPE-1" found.')])
        self.assertIn(('Employee ID', 'No employee with Employee ID "NOBODY" found.'), errs[7])
        self.assertIn(repeated, errs[7])
        self.assertIn(repeated, errs[8])
        self.assertEqual(errs[9], [('Employee ID', 'This employee is inactive.')])
        self.assertEqual(errs[10], [('Employee ID', 'This employee has not completed onboarding yet.')])
        self.assertIn(('Employee ID', 'Employee ID "DUPX" matches more than one employee; it can\'t be used to assign.'),
                      errs[11])                                                   # ambiguous (DUPX / dupx)
        self.assertIn(('Asset Tag', '"ASG-REP" is not available — current status: Under Repair.'), errs[11])
        self.assertEqual(self._import(_csv(rows)).data['data']['summary']['assigned'], 0)
        self.assertFalse(AssetAssignment.objects.exists())

    def test_exact_case_employee_id_still_ambiguous_when_two_match(self):
        _login(self.client, self.admin.email)
        data = self._import(_csv([_row('ASG-A1', 'dupx')]), dry_run=True).data['data']
        self.assertEqual(data['rows'][0]['status'], 'error')
        self.assertIn('matches more than one employee', data['rows'][0]['errors'][0]['message'])

    # ── already assigned, retry, partial success, error report ───────────
    def test_already_assigned_partial_success_report_and_retry(self):
        assign_asset(asset=self.a3, employee=self.emp_a1, assigned_date=datetime.date(2026, 1, 1),
                     condition_at_assignment='good', assigned_by=self.admin)
        assign_asset(asset=self.a4, employee=self.emp_a2, assigned_date=datetime.date(2026, 1, 1),
                     condition_at_assignment='good', assigned_by=self.admin)
        _login(self.client, self.admin.email)
        rows = [
            _row('ASG-A1', 'EMPA1'),            # assigned
            _row('ASG-A3', 'EMPA1'),            # already with this employee
            _row('ASG-A4', 'EMPA1'),            # already with someone else — never reassigned
            _row('ASG-A2', 'NOBODY'),           # failed
        ]
        preview = self._import(_csv(rows), dry_run=True).data['data']
        self.assertEqual(preview['counts'], {'ready': 1, 'already_assigned': 2, 'error': 1})
        errs = self._errors(preview)
        self.assertEqual(errs[3], [('Asset Tag', 'Already assigned to this employee.')])
        self.assertEqual(errs[4], [('Asset Tag', 'Already assigned to another employee. Return it first to reassign.')])

        resp = self._import(_csv(rows))
        self.assertEqual(resp.data['data']['summary'],
                         {'total': 4, 'assigned': 1, 'already_assigned': 2, 'failed': 1, 'skipped': 0})
        self.assertEqual(AssetAssignment.objects.get(asset=self.a4, status='assigned').employee, self.emp_a2)
        report = base64.b64decode(resp.data['data']['error_report_csv']).decode('utf-8-sig')
        lines = list(csv.reader(io.StringIO(report)))
        self.assertEqual(lines[0], ['Row', 'Asset Tag', 'Employee ID', 'Status', 'Column', 'Reason'])
        self.assertEqual([line[:4] for line in lines[1:]], [
            ['3', 'ASG-A3', 'EMPA1', 'already assigned'], ['4', 'ASG-A4', 'EMPA1', 'already assigned'],
            ['5', 'ASG-A2', 'NOBODY', 'error']])
        self.assertNotIn('Traceback', report)

        before = AssetAssignment.objects.count()
        s = self._import(_csv(rows)).data['data']['summary']               # retried: nothing created twice
        self.assertEqual((s['assigned'], s['already_assigned'], s['failed']), (0, 3, 1))
        self.assertEqual(AssetAssignment.objects.count(), before)

    def test_report_cells_cannot_become_formulas(self):
        _login(self.client, self.admin.email)
        resp = self._import(_csv([_row('=HYPERLINK("x")', '+EMP')]), dry_run=True)
        report = base64.b64decode(resp.data['data']['error_report_csv']).decode('utf-8-sig')
        for line in list(csv.reader(io.StringIO(report)))[1:]:
            for cell in line:
                self.assertNotIn(cell[:1], ('=', '+', '-', '@'))

    # ── permissions and branch scope ─────────────────────────────────────
    def test_requires_assets_edit(self):
        _login(self.client, self.creator.email)                     # assets.create only
        self.assertEqual(self._import(_csv([_row('ASG-A1', 'EMPA1')])).status_code, 403)
        self.assertEqual(self.client.get(reverse('asset-assign-import-sample')).status_code, 403)
        self.assertFalse(AssetAssignment.objects.exists())
        self.client.cookies.clear()
        self.assertIn(self._import(_csv([_row('ASG-A1', 'EMPA1')])).status_code, (401, 403))

    def test_branch_user_scope(self):
        _login(self.client, self.branch_user.email)
        rows = [
            _row('ASG-A1', 'EMPA1'),            # own branch asset + employee
            _row('ASG-B1', 'EMPA2'),            # other branch's asset: looks unknown
            _row('ASG-A2', 'EMPB1'),            # other branch's employee: looks unknown
        ]
        data = self._import(_csv(rows)).data['data']
        self.assertEqual(data['summary']['assigned'], 1)
        errs = self._errors(data)
        self.assertEqual(errs[3], [('Asset Tag', 'No asset with Asset Tag "ASG-B1" found in your branch.')])
        self.assertEqual(errs[4], [('Employee ID', 'No employee with Employee ID "EMPB1" found.')])
        self.assertFalse(AssetAssignment.objects.filter(asset__in=[self.a2, self.b1]).exists())

    def test_manager_limited_to_direct_reports(self):
        _login(self.client, self.manager.email)
        data = self._import(_csv([_row('ASG-A1', 'EMPA1'), _row('ASG-A2', 'EMPA2')])).data['data']
        self.assertEqual(data['summary']['assigned'], 1)                  # EMPA1 reports to the manager
        self.assertEqual(self._errors(data)[3], [('Employee ID', 'No employee with Employee ID "EMPA2" found.')])

    # ── concurrency: state changes between preview and confirm ──────────
    def test_asset_taken_after_planning_is_not_reassigned(self):
        rows = parse_file(_csv([_row('ASG-A1', 'EMPA1'), _row('ASG-A2', 'EMPA1')]))
        plan = plan_assignments(rows, user=self.admin, is_admin=True, user_branch=None)
        self.assertEqual(plan['counts']['ready'], 2)
        # another request assigns A1 and sends A2 to repair before this plan is applied
        assign_asset(asset=self.a1, employee=self.emp_a2, assigned_date=datetime.date(2026, 1, 1),
                     condition_at_assignment='good', assigned_by=self.admin)
        Asset.objects.filter(pk=self.a2.pk).update(status=Asset.STATUS_UNDER_REPAIR)
        result = apply_assignments(plan, user=self.admin)
        self.assertEqual([r['status'] for r in result['rows']], ['already_assigned', 'error'])
        self.assertEqual(result['rows'][1]['errors'][0]['message'],
                         '"ASG-A2" is not available — current status: Under Repair.')
        self.assertEqual(AssetAssignment.objects.get(asset=self.a1).employee, self.emp_a2)
        self.assertFalse(AssetAssignment.objects.filter(asset=self.a2).exists())

    def test_unique_constraint_race_reported_not_raised(self):
        rows = parse_file(_csv([_row('ASG-A1', 'EMPA1')]))
        plan = plan_assignments(rows, user=self.admin, is_admin=True, user_branch=None)
        with mock.patch('apps.assets.models.AssetAssignment.objects.create', side_effect=IntegrityError('dup')):
            result = apply_assignments(plan, user=self.admin)
        self.assertEqual(result['rows'][0]['status'], 'already_assigned')
        self.assertNotIn('dup', str(result['rows'][0]['errors']))
        self.a1.refresh_from_db()
        self.assertEqual(self.a1.status, Asset.STATUS_AVAILABLE)                # rolled back


class SingleAssignRegressionTests(TestCase):
    """Single Assign Asset / Return / history behave exactly as before the shared-service refactor."""

    def setUp(self):
        cache.clear()
        self.branch = _make_branch('SingleAsgBranch')
        role = make_role('single_asg_hr', permission_codenames=['assets.view', 'assets.edit', 'settings.edit'])
        self.hr = make_user('single.hr@test.com', role=role, employee_id='SAHR', full_name='Single HR',
                            branch=self.branch.branch_name)
        self.employee = make_user('single.emp@test.com', employee_id='SAEMP', full_name='Single Emp',
                                  branch=self.branch.branch_name, onboarding_status=User.ONBOARDING_COMPLETE)
        self.asset = Asset.objects.create(asset_tag='SA-1', asset_name='Phone', category='IT', asset_type='Phone',
                                          branch=self.branch)
        self.api = APIClient()
        _login(self.api, self.hr.email)

    def _assign(self, employee=None, **over):
        body = {'asset': str(self.asset.pk), 'assigned_date': '2026-01-15', 'condition_at_assignment': 'fair',
                'expected_return_date': '2026-06-30', 'remarks': 'Field visit'}
        body.update(over)
        return self.api.post(reverse('asset-assign', kwargs={'employee_id': (employee or self.employee).pk}),
                             body, format='json')

    def test_single_assign_response_history_and_audit_unchanged(self):
        resp = self._assign()
        self.assertEqual(resp.status_code, 201, resp.data)
        self.assertEqual(resp.data['message'], 'Asset assigned successfully.')
        self.assertEqual(resp.data['data']['asset_tag'], 'SA-1')
        asg = AssetAssignment.objects.get(asset=self.asset)
        self.assertEqual((asg.condition_at_assignment, str(asg.expected_return_date), asg.assign_remarks, asg.assigned_by),
                         ('fair', '2026-06-30', 'Field visit', self.hr))
        self.asset.refresh_from_db()
        self.assertEqual((self.asset.status, self.asset.condition), (Asset.STATUS_ASSIGNED, 'fair'))
        log = AuditLog.objects.get(action='asset_assigned')
        self.assertEqual(log.changes, {'asset_tag': 'SA-1', 'employee': 'Single Emp'})     # no "source" key
        self.assertEqual((log.object_id, log.branch), (str(asg.pk), self.branch.branch_name))

    def test_single_assign_unavailable_under_lock_is_409_with_same_message(self):
        # the serializer check passes, then the row-locked re-check catches the change
        with mock.patch('apps.assets.serializers.AssignAssetSerializer.validate_asset', side_effect=lambda v: v):
            Asset.objects.filter(pk=self.asset.pk).update(status=Asset.STATUS_LOST)
            resp = self._assign()
        self.assertEqual(resp.status_code, 409)
        self.assertEqual(resp.data['message'], '"SA-1" is not available — current status: Lost.')

    def test_single_assign_integrity_race_is_409_with_same_message(self):
        with mock.patch('apps.assets.models.AssetAssignment.objects.create', side_effect=IntegrityError('x')):
            resp = self._assign()
        self.assertEqual(resp.status_code, 409)
        self.assertEqual(resp.data['message'],
                         '"SA-1" was just assigned to someone else — please refresh and try again.')

    def test_single_assign_still_allows_pending_employee(self):
        # The onboarding/active checks are bulk-only; single assign is unchanged.
        pending = make_user('single.pend@test.com', employee_id='SAPEND', branch=self.branch.branch_name)
        self.assertEqual(self._assign(employee=pending).status_code, 201)

    def test_return_then_bulk_reassign_keeps_history(self):
        assignment_id = self._assign().data['data']['id']
        resp = self.api.post(reverse('asset-return', kwargs={'assignment_id': assignment_id}), {
            'return_date': '2026-02-01', 'return_condition': 'good', 'return_reason': 'Project ended',
        }, format='json')
        self.assertEqual(resp.status_code, 200, resp.data)
        upload = _csv([_row('SA-1', 'SAEMP', **{'Assigned Date': '2026-03-01'})])
        resp = self.api.post(reverse('asset-assign-import'), {'file': upload}, format='multipart')
        self.assertEqual(resp.data['data']['summary']['assigned'], 1, resp.data)
        history = self.api.get(reverse('employee-assets', kwargs={'employee_id': self.employee.pk})).data['data']
        self.assertEqual((len(history['current']), len(history['history'])), (1, 1))
        self.assertEqual(history['history'][0]['return_reason'], 'Project ended')
